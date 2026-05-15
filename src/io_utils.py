"""Input utilities for experiment data files."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

import numpy as np
from openpyxl import load_workbook


def _numeric_column_values(rows: Iterable[tuple], column_index: int) -> list[float]:
    values: list[float] = []
    for row in rows:
        if column_index >= len(row):
            continue
        value = row[column_index]
        if isinstance(value, (int, float)):
            values.append(float(value))
    return values


def read_bit_column(path: str | Path, sheet_name: str) -> np.ndarray:
    """Read a one-column bit sequence from an xlsx sheet."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[sheet_name]
    bits = _numeric_column_values(sheet.iter_rows(values_only=True), 0)
    bits_array = np.asarray(bits, dtype=int)
    if not np.isin(bits_array, [0, 1]).all():
        raise ValueError(f"{sheet_name} contains values other than 0 and 1.")
    return bits_array


def read_complex_signal(path: str | Path, sheet_name: str) -> np.ndarray:
    """Read u_I and u_Q columns from an xlsx sheet as complex symbols."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    i_values = _numeric_column_values(rows, 0)
    q_values = _numeric_column_values(rows, 1)
    if len(i_values) != len(q_values):
        raise ValueError(f"{sheet_name} has different u_I and u_Q lengths.")
    return np.asarray(i_values, dtype=float) + 1j * np.asarray(q_values, dtype=float)


def split_frame(symbols: np.ndarray, preamble_symbols: int = 16) -> tuple[np.ndarray, np.ndarray]:
    """Split a frame into preamble and data symbols."""
    symbols_array = np.asarray(symbols, dtype=complex).reshape(-1)
    if symbols_array.size <= preamble_symbols:
        raise ValueError("Frame does not contain data symbols after the preamble.")
    return symbols_array[:preamble_symbols], symbols_array[preamble_symbols:]


def describe_xlsx(path: str | Path) -> str:
    """Return a compact text description of workbook sheets and first rows."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    lines = [f"file: {path}"]
    for sheet in workbook.worksheets:
        lines.append(f"sheet: {sheet.title}, rows={sheet.max_row}, cols={sheet.max_column}")
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8), values_only=True):
            lines.append(f"  {row}")
    return "\n".join(lines)
